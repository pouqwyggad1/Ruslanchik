import os
import re
import sys
import time

import certifi
import openpyxl
import undetected_chromedriver as uc
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait

os.environ.setdefault("SSL_CERT_FILE", certifi.where())


def start_browser():
    """Запускает Chrome и возвращает объект для управления браузером."""
    print("Запуск Chrome")

    options = uc.ChromeOptions()
    options.add_argument("--start-maximized")
    options.add_argument("--disable-notifications")

    driver = uc.Chrome(options=options)
    return driver


def close_popups(driver):
    """Закрывает всплывающие окна, если они появились."""

    # На сайте могут появиться разные кнопки закрытия.
    # Для каждой кнопки указан свой XPath.
    popups = [
        "//button[contains(text(), 'Все верно')]",
        "//button[contains(text(), 'Понятно')]",
        "//button[@aria-label='Закрыть']",
    ]

    for xpath in popups:
        try:
            buttons = driver.find_elements(By.XPATH, xpath)
            if buttons:
                buttons[0].click()
                time.sleep(1)
        except Exception:
            # Если окна нет или оно уже закрылось, продолжаем работу.
            pass


def search_products(driver, search_text):
    """Открывает DNS и выполняет поиск товаров."""
    print(f"Поиск товаров по запросу: {search_text}")

    driver.get("https://www.dns-shop.ru/")
    close_popups(driver)

    # Ждём, пока строка поиска станет доступна для ввода.
    wait = WebDriverWait(driver, 20)
    search_input = wait.until(
        EC.element_to_be_clickable(
            (By.XPATH, "//input[@type='search' or contains(@placeholder, 'Поиск')]")
        )
    )

    search_input.clear()
    search_input.send_keys(search_text)
    search_input.send_keys(Keys.ENTER)

    # Ждём появления хотя бы одной карточки товара.
    wait.until(
        EC.presence_of_element_located(
            (By.CSS_SELECTOR, "[data-id='product'], .catalog-product")
        )
    )
    time.sleep(2)
    close_popups(driver)


def sort_products(driver):
    """Выбирает сортировку по количеству отзывов."""
    print("Сортировка товаров по количеству отзывов")

    wait = WebDriverWait(driver, 15)

    # Открываем список вариантов сортировки
    sort_button = wait.until(
        EC.element_to_be_clickable(
            (
                By.XPATH,
                "//*[contains(@class, 'top-filter__label')]/.. "
                "| //span[contains(text(), 'Сортировка')]/..",
            )
        )
    )
    sort_button.click()

    # Выбираем нужный пункт
    reviews_button = wait.until(
        EC.element_to_be_clickable(
            (
                By.XPATH,
                "//label[contains(., 'По количеству отзывов')] "
                "| //span[contains(., 'По количеству отзывов')]",
            )
        )
    )
    reviews_button.click()

    # Пауза нужна, чтобы каталог успел обновиться
    time.sleep(4)


def get_number(text):
    """Оставляет в строке только цифры и возвращает целое число."""
    digits = re.sub(r"[^0-9]", "", text)
    return int(digits) if digits else 0


def collect_products(driver):
    """Собирает данные обо всех товарах с первой страницы."""
    print("Сбор названий, цен и рейтингов товаров")

    cards = driver.find_elements(
        By.CSS_SELECTOR, "[data-id='product'], .catalog-product"
    )
    products = []

    for card in cards:
        try:
            # Получаем название товара
            title = card.find_element(
                By.CSS_SELECTOR,
                "a.catalog-product__name, [data-role='product-title']",
            ).text.strip()

            # Из строки "5 999 ₽" получаем число 5999
            try:
                price_text = card.find_element(
                    By.CSS_SELECTOR,
                    ".product-buy__price, [class*='price__current']",
                ).text
                price = get_number(price_text)
            except Exception:
                price = 0

            # Запятую в рейтинге заменяем точкой для преобразования в float
            try:
                rating_text = card.find_element(
                    By.CSS_SELECTOR,
                    "a.catalog-product__rating, [data-rating], [class*='rating']",
                ).text
                result = re.search(r"\d+(?:[.,]\d+)?", rating_text)
                rating = float(result.group().replace(",", ".")) if result else 0
            except Exception:
                rating = 0

            # Ищем кнопку добавления в избранное внутри карточки
            favorite_button = None
            buttons = card.find_elements(
                By.CSS_SELECTOR,
                "button[class*='wishlist'], button[class*='favorite'], "
                "button[class*='like'], button.button-ui_white",
            )

            for button in buttons:
                name = (
                    (button.get_attribute("title") or "")
                    + (button.get_attribute("aria-label") or "")
                    + (button.get_attribute("class") or "")
                ).lower()

                if "избран" in name or "wishlist" in name or "like" in name:
                    favorite_button = button
                    break

            # Все данные одного товара хранятся в словаре
            products.append(
                {
                    "title": title,
                    "price": price,
                    "rating": rating,
                    "favorite_button": favorite_button,
                    "status": "-",
                }
            )

        except Exception as error:
            # Ошибка одной карточки не останавливает сбор остальных
            print("Не удалось прочитать одну карточку:", error)

    print("Найдено товаров:", len(products))
    return products


def add_cheapest_to_favorites(driver, products):
    """Находит самый дешёвый товар и добавляет его в избранное."""
    print("Добавление самого дешёвого товара в избранное")

    # Нулевую цену не учитываем: она означает, что цена не распозналась
    products_with_price = [product for product in products if product["price"] > 0]

    if not products_with_price:
        raise RuntimeError("Не удалось получить цены товаров")

    # key указывает, что товары нужно сравнивать по полю price
    cheapest = min(products_with_price, key=lambda product: product["price"])
    button = cheapest["favorite_button"]

    print("Самый дешёвый товар:", cheapest["title"])
    print("Цена:", cheapest["price"], "руб.")

    if button is None:
        raise RuntimeError("Не найдена кнопка добавления в избранное")

    driver.execute_script(
        "arguments[0].scrollIntoView({block: 'center'});", button
    )

    try:
        button.click()
    except Exception:
        # Запасной способ нажатия через JavaScript
        driver.execute_script("arguments[0].click();", button)

    # Меняем статус только после нажатия кнопки
    cheapest["status"] = "+"
    time.sleep(2)
    close_popups(driver)


def create_excel(products, filename):
    """Создаёт и оформляет Excel-отчёт."""
    print("Создание Excel")

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Товары"

    # Сортируем товары по названию без учёта регистра букв
    products.sort(key=lambda product: product["title"].lower())

    # Записываем заголовки и данные
    sheet.append(["Название", "Цена", "Рейтинг", "Статус"])

    for product in products:
        sheet.append(
            [
                product["title"],
                product["price"],
                product["rating"],
                product["status"],
            ]
        )

    last_row = len(products) + 1

    # Создаём стили, которые будем применять к ячейкам
    blue_fill = PatternFill("solid", fgColor="366092")
    white_font = Font(color="FFFFFF", bold=True)
    gray_line = Side(style="thin", color="D9D9D9")
    border = Border(left=gray_line, right=gray_line, top=gray_line, bottom=gray_line)

    # Оформляем шапку таблицы
    for cell in sheet[1]:
        cell.fill = blue_fill
        cell.font = white_font
        cell.alignment = Alignment(horizontal="center")

    # Оформляем строки с товарами
    for row in range(2, last_row + 1):
        for column in range(1, 5):
            sheet.cell(row, column).border = border

        sheet.cell(row, 2).number_format = '#,##0 "₽"'
        sheet.cell(row, 3).number_format = "0.00"
        sheet.cell(row, 3).alignment = Alignment(horizontal="center")
        sheet.cell(row, 4).alignment = Alignment(horizontal="center")

    # Добавляем фильтр на всю таблицу.
    sheet.auto_filter.ref = f"A1:D{last_row}"

    # Под таблицей записываем формулу минимальной цены
    formula_row = last_row + 1
    sheet.cell(formula_row, 1, "Минимальная цена").font = Font(bold=True)
    sheet.cell(formula_row, 2, f"=MIN(B2:B{last_row})").font = Font(bold=True)
    sheet.cell(formula_row, 2).number_format = '#,##0 "₽"'

    # Цвета для условного форматирования
    green_fill = PatternFill("solid", fgColor="C6EFCE")
    red_fill = PatternFill("solid", fgColor="FFC7CE")

    # Рейтинг 4.8 и выше будет зелёным, остальные - красными
    sheet.conditional_formatting.add(
        f"C2:C{last_row}",
        CellIsRule(operator="greaterThanOrEqual", formula=["4.8"], fill=green_fill),
    )
    sheet.conditional_formatting.add(
        f"C2:C{last_row}",
        CellIsRule(operator="lessThan", formula=["4.8"], fill=red_fill),
    )

    # Статус "+" - зелёным, а статус "-" - красным.
    sheet.conditional_formatting.add(
        f"D2:D{last_row}",
        CellIsRule(operator="equal", formula=['"+"'], fill=green_fill),
    )
    sheet.conditional_formatting.add(
        f"D2:D{last_row}",
        CellIsRule(operator="equal", formula=['"-"'], fill=red_fill),
    )

    # Создаём выпадающий список, в котором можно выбрать только + или -
    status_list = DataValidation(type="list", formula1='"+,-"')
    sheet.add_data_validation(status_list)
    status_list.add(f"D2:D{last_row}")

    # Настраиваем ширину столбцов
    sheet.column_dimensions["A"].width = 55
    sheet.column_dimensions["B"].width = 15
    sheet.column_dimensions["C"].width = 15
    sheet.column_dimensions["D"].width = 15

    workbook.save(filename)
    print("Файл сохранён:", filename)


def open_favorites(driver):
    """Открывает страницу избранных товаров."""
    print("Открываем избранное")

    # Прямой переход
    driver.get("https://www.dns-shop.ru/profile/wishlist/")
    time.sleep(3)
    close_popups(driver)


def main():
    """Главная функция: по очереди вызывает все этапы программы."""
    # Если запрос не указан, программа ищет мыши
    # sys.argv содержит слова, написанные после имени файла
    search_text = " ".join(sys.argv[1:]) if len(sys.argv) > 1 else "мышь"
    driver = None

    try:
        driver = start_browser()
        search_products(driver, search_text)
        sort_products(driver)

        products = collect_products(driver)
        if not products:
            raise RuntimeError("На странице не найдено ни одного товара")

        add_cheapest_to_favorites(driver, products)
        create_excel(products, "products.xlsx")
        open_favorites(driver)

        input("Готово. Нажмите Enter, чтобы закрыть браузер")

    except Exception as error:
        print("Ошибка:", error)

    finally:
        # finally выполняется и при успешной работе, и при ошибке
        if driver is not None:
            driver.quit()

if __name__ == "__main__":
    main()
